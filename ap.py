import json
import os
import sqlite3
import threading
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
DB_FILE = ROOT / "infinix_shopee.db"
EXCEL_FILE = ROOT / "infinix_shopee_database.xlsx"
SAVE_LOCK = threading.Lock()

# Browser data keys mapped to readable Excel sheet names.
COLLECTIONS = {
    "infinix-accounts": "Accounts",
    "infinix-products": "Products",
    "infinix-orders": "Orders",
    "infinix-plus-members": "Memberships",
    "infinix-plus-requests": "Membership Requests",
    "infinix-vouchers": "Vouchers",
    "infinix-activity-log": "Recent Activities",
    "infinix-delivery-locations": "Delivery Addresses",
    "infinix-promotion-settings": "Promotions",
    "infinix-plus-plan-prices": "Plus Prices",
    "shoply-cart": "Carts",
    "shoply-wish": "Wishlists",
}


def connect_db():
    connection = sqlite3.connect(DB_FILE)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS app_state (
            storage_key TEXT PRIMARY KEY,
            value_json TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
    """)
    return connection


def read_state():
    with connect_db() as connection:
        rows = connection.execute(
            "SELECT storage_key, value_json FROM app_state"
        ).fetchall()
    return {key: value for key, value in rows}


def excel_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if value is None:
        return ""
    return value


def records_from(value):
    """Turn list or object data into rows for the Excel collection sheets."""
    if isinstance(value, list):
        return [
            item if isinstance(item, dict) else {"value": item}
            for item in value
        ]

    if isinstance(value, dict):
        rows = []
        for record_id, item in value.items():
            if isinstance(item, dict):
                rows.append({"record_id": record_id, **item})
            else:
                rows.append({"record_id": record_id, "value": item})
        return rows

    return [{"value": value}]


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E4B")
    for column in sheet.columns:
        letter = column[0].column_letter
        longest = max(
            (len(str(cell.value or "")) for cell in column),
            default=12,
        )
        sheet.column_dimensions[letter].width = min(max(longest + 2, 14), 42)


def export_to_excel(state):
    workbook = Workbook()
    all_data = workbook.active
    all_data.title = "All App Data"
    all_data.append(["Browser Data Key", "Saved Data", "Updated At"])

    updated_at = datetime.now(timezone.utc).isoformat()

    for key, raw_value in sorted(state.items()):
        all_data.append([key, raw_value, updated_at])

    format_sheet(all_data)

    for key, sheet_name in COLLECTIONS.items():
        raw_value = state.get(key)
        if raw_value is None:
            continue

        try:
            parsed = json.loads(raw_value)
        except (TypeError, json.JSONDecodeError):
            parsed = {"value": raw_value}

        rows = records_from(parsed)
        headers = sorted({field for row in rows for field in row.keys()})
        if not headers:
            headers = ["value"]

        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)

        for row in rows:
            sheet.append([excel_value(row.get(header)) for header in headers])

        format_sheet(sheet)

    temporary_file = EXCEL_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary_file)
    os.replace(temporary_file, EXCEL_FILE)


def save_state(state):
    timestamp = datetime.now(timezone.utc).isoformat()

    with SAVE_LOCK:
        with connect_db() as connection:
            connection.execute("DELETE FROM app_state")
            connection.executemany(
                """
                INSERT INTO app_state (storage_key, value_json, updated_at)
                VALUES (?, ?, ?)
                """,
                [
                    (key, value, timestamp)
                    for key, value in state.items()
                ],
            )

        export_to_excel(state)


class ShopHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, status, data):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if urlsplit(self.path).path == "/api/state":
            self.send_json(200, read_state())
            return
        super().do_GET()

    def do_POST(self):
        if urlsplit(self.path).path != "/api/state":
            self.send_error(404)
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length) or b"{}")

            if not isinstance(payload, dict):
                raise ValueError("Expected a JSON object")

            # localStorage values are strings. Keep the same format so the
            # existing website can load them again.
            state = {
                str(key): str(value)
                for key, value in payload.items()
            }

            save_state(state)
            self.send_json(200, {"saved": True})
        except Exception as error:
            self.send_json(400, {"saved": False, "error": str(error)})

    do_PUT = do_POST


if __name__ == "__main__":
    connect_db().close()
    if not EXCEL_FILE.exists():
        export_to_excel(read_state())

    server = ThreadingHTTPServer(("127.0.0.1", 8000), ShopHandler)
    print("Infinix Shopee is running at http://127.0.0.1:8000")
    print(f"Excel data file: {EXCEL_FILE}")
    server.serve_forever()
