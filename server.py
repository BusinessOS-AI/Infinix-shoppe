"""Infinix Shopee application server.

Authentication is handled here, never in browser storage. Set SESSION_SECRET,
ADMIN_USERNAME, and ADMIN_PASSWORD as Render environment variables before use.
"""
import hashlib
import hmac
import json
import os
import secrets
import sqlite3
import threading
import time
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent
DB_FILE = Path(os.environ.get("INFINIX_DB_PATH", str(ROOT / "infinix_shopee.db")))
EXCEL_FILE = Path(os.environ.get("INFINIX_EXCEL_PATH", str(ROOT / "infinix_shopee_database.xlsx")))
SESSION_SECRET = os.environ.get("SESSION_SECRET", "")
ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "").strip()
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
SESSION_SECONDS = 60 * 60 * 24 * 7
PBKDF2_ROUNDS = 310_000
SAVE_LOCK = threading.RLock()
DB_FILE.parent.mkdir(parents=True, exist_ok=True)
EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)

COLLECTIONS = {
    "infinix-accounts": "Accounts", "infinix-products": "Products",
    "infinix-orders": "Orders", "infinix-plus-members": "Memberships",
    "infinix-plus-requests": "Membership Requests", "infinix-vouchers": "Vouchers",
    "infinix-activity-log": "Recent Activities", "infinix-delivery-locations": "Delivery Addresses",
    "infinix-promotion-settings": "Promotions", "infinix-plus-plan-prices": "Plus Prices",
    "shoply-cart": "Carts", "shoply-wish": "Wishlists",
}
PUBLIC_KEYS = {"infinix-products", "infinix-vouchers", "infinix-promotion-settings", "infinix-plus-plan-prices"}
ADMIN_KEYS = {"infinix-vouchers", "infinix-promotion-settings", "infinix-plus-plan-prices", "infinix-plus-members", "infinix-activity-log"}
PRIVATE_KEYS = {"shoply-cart", "shoply-wish", "infinix-delivery-locations", "infinix-applied-voucher", "infinix-theme", "infinix-country", "infinix-user", "infinix-role", "infinix-catalog-version"}
SAVED_PRIVATE_KEYS = PRIVATE_KEYS - {"infinix-user", "infinix-role", "infinix-catalog-version"}


def connect_db():
    conn = sqlite3.connect(DB_FILE, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("CREATE TABLE IF NOT EXISTS app_state (storage_key TEXT PRIMARY KEY, value_json TEXT NOT NULL, updated_at TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS user_state (user_id INTEGER NOT NULL, storage_key TEXT NOT NULL, value_json TEXT NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(user_id, storage_key))")
    conn.execute("CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT NOT NULL UNIQUE COLLATE NOCASE, email TEXT NOT NULL UNIQUE COLLATE NOCASE, name TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('customer','shopkeeper','admin')), salt TEXT NOT NULL, password_hash TEXT NOT NULL, created_at TEXT NOT NULL)")
    conn.execute("CREATE TABLE IF NOT EXISTS sessions (token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL, expires_at INTEGER NOT NULL)")
    # Migrate legacy server-side accounts to password hashes before removing the
    # old JSON collection, which may contain plain-text passwords.
    legacy = conn.execute("SELECT value_json FROM app_state WHERE storage_key='infinix-accounts'").fetchone()
    if legacy:
        try:
            old_accounts = json.loads(legacy["value_json"])
        except (TypeError, json.JSONDecodeError):
            old_accounts = []
        if isinstance(old_accounts, list):
            for account in old_accounts:
                if not isinstance(account, dict):
                    continue
                username = str(account.get("username", "")).strip()
                email = str(account.get("email", "")).strip().lower()
                name = str(account.get("name", username)).strip() or username
                password = str(account.get("password", ""))
                role = str(account.get("role", "customer"))
                if (role not in {"customer", "shopkeeper"} or len(username) < 3
                        or len(username) > 32 or not username.replace("_", "").isalnum()
                        or "@" not in email or len(password) < 8):
                    continue
                salt, digest = password_digest(password)
                try:
                    conn.execute(
                        "INSERT INTO users(username,email,name,role,salt,password_hash,created_at) VALUES(?,?,?,?,?,?,?)",
                        (username, email, name[:100], role, salt, digest, now_iso()),
                    )
                except sqlite3.IntegrityError:
                    pass
    conn.execute("DELETE FROM app_state WHERE storage_key='infinix-accounts'")
    conn.commit()
    return conn


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def password_digest(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ROUNDS)
    return salt.hex(), digest.hex()


def verify_password(password, salt_hex, digest_hex):
    try:
        _, candidate = password_digest(password, bytes.fromhex(salt_hex))
        return hmac.compare_digest(candidate, digest_hex)
    except (ValueError, TypeError):
        return False


def token_digest(token):
    secret = (SESSION_SECRET or "local-development-only").encode("utf-8")
    return hmac.new(secret, token.encode("utf-8"), hashlib.sha256).hexdigest()


def public_user(row):
    return {"id": row["id"], "username": row["username"], "email": row["email"], "name": row["name"], "role": row["role"], "createdAt": row["created_at"]}


def get_user_by_session(token):
    if not token:
        return None
    hashed = token_digest(token)
    with connect_db() as conn:
        row = conn.execute("SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=? AND s.expires_at>?", (hashed, int(time.time()))).fetchone()
    return public_user(row) if row else None


def issue_session(user_id):
    token = secrets.token_urlsafe(36)
    with connect_db() as conn:
        conn.execute("INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)", (token_digest(token), user_id, int(time.time()) + SESSION_SECONDS))
        conn.execute("DELETE FROM sessions WHERE expires_at<=?", (int(time.time()),))
    return token


def read_global():
    with connect_db() as conn:
        return {r["storage_key"]: r["value_json"] for r in conn.execute("SELECT storage_key,value_json FROM app_state")}


def parse_json(value, fallback):
    try:
        return json.loads(value) if value is not None else fallback
    except (TypeError, json.JSONDecodeError):
        return fallback


def account_list():
    with connect_db() as conn:
        return [public_user(row) for row in conn.execute("SELECT * FROM users ORDER BY created_at")]


def get_visible_state(user):
    global_state = read_global()
    result = {}
    if user:
        with connect_db() as conn:
            result.update({r["storage_key"]: r["value_json"] for r in conn.execute("SELECT storage_key,value_json FROM user_state WHERE user_id=?", (user["id"],))})
    # A visitor can see only the public catalog and offer configuration.
    if not user:
        allowed = PUBLIC_KEYS
    elif user["role"] == "admin":
        allowed = set(global_state) - PRIVATE_KEYS
    else:
        allowed = PUBLIC_KEYS | {"infinix-orders"}
        if user["role"] == "customer":
            allowed |= {"infinix-plus-members", "infinix-plus-requests"}
    for key in allowed:
        if key in global_state and key not in {"infinix-accounts"}:
            result[key] = global_state[key]
    if not user or user["role"] != "admin":
        offers = parse_json(result.get("infinix-vouchers"), [])
        result["infinix-vouchers"] = json.dumps([v for v in offers if isinstance(v, dict) and v.get("active")], ensure_ascii=False)
    if not user:
        return result
    role = user["role"]
    ident = user["username"].strip().lower()
    email = user["email"].strip().lower()
    if role == "admin":
        result["infinix-accounts"] = json.dumps(account_list(), ensure_ascii=False)
        return result
    # Orders and membership records are filtered server-side by account identity.
    orders = parse_json(global_state.get("infinix-orders"), [])
    if role == "customer":
        orders = [o for o in orders if str(o.get("customer", "")).lower() in {ident, email}]
        result["infinix-orders"] = json.dumps(orders, ensure_ascii=False)
        members = parse_json(global_state.get("infinix-plus-members"), {})
        requests = parse_json(global_state.get("infinix-plus-requests"), [])
        own_members = {k: v for k, v in members.items() if str(k).lower() in {ident, email}}
        own_requests = [r for r in requests if str(r.get("username", "")).lower() in {ident, email} or str(r.get("email", "")).lower() in {ident, email}]
        result["infinix-plus-members"] = json.dumps(own_members, ensure_ascii=False)
        result["infinix-plus-requests"] = json.dumps(own_requests, ensure_ascii=False)
    elif role == "shopkeeper":
        orders = [o for o in orders if any(str(i.get("seller", "")).lower() in {ident, email} for i in o.get("items", []))]
        result["infinix-orders"] = json.dumps(orders, ensure_ascii=False)
    result.pop("infinix-activity-log", None)
    result.pop("infinix-plus-members", None) if role != "customer" else None
    result.pop("infinix-plus-requests", None) if role != "customer" else None
    result.pop("infinix-accounts", None)
    return result


def records_from(value):
    if isinstance(value, list):
        return [item if isinstance(item, dict) else {"value": item} for item in value]
    if isinstance(value, dict):
        return [{"record_id": key, **item} if isinstance(item, dict) else {"record_id": key, "value": item} for key, item in value.items()]
    return [{"value": value}]


def style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E4B")
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max((len(str(cell.value or "")) for cell in column), default=12)
        sheet.column_dimensions[letter].width = min(max(width + 2, 14), 42)


def export_excel(state):
    book = Workbook()
    all_data = book.active
    all_data.title = "All App Data"
    all_data.append(["Data Key", "Saved Data", "Updated At"])
    timestamp = now_iso()
    for key, value in sorted(state.items()):
        all_data.append([key, value, timestamp])
    style_sheet(all_data)
    for key, title in COLLECTIONS.items():
        if key == "infinix-accounts" or key not in state:
            continue
        rows = records_from(parse_json(state[key], {"value": state[key]}))
        headers = sorted({name for row in rows for name in row}) or ["value"]
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([json.dumps(row.get(h), ensure_ascii=False) if isinstance(row.get(h), (dict, list)) else row.get(h, "") for h in headers])
        style_sheet(sheet)
    with connect_db() as conn:
        safe_users = conn.execute("SELECT username,email,name,role,created_at FROM users ORDER BY created_at").fetchall()
        private_rows = conn.execute("SELECT u.username,s.storage_key,s.value_json,s.updated_at FROM user_state s JOIN users u ON u.id=s.user_id ORDER BY u.username,s.storage_key").fetchall()
    users_sheet = book.create_sheet("Users")
    users_sheet.append(["Username", "Email", "Name", "Role", "Created At"])
    for row in safe_users:
        users_sheet.append(list(row))
    style_sheet(users_sheet)
    private_sheet = book.create_sheet("User Private Data")
    private_sheet.append(["Username", "Data Key", "Saved Data", "Updated At"])
    for row in private_rows:
        private_sheet.append(list(row))
    style_sheet(private_sheet)
    temp = EXCEL_FILE.with_suffix(".tmp.xlsx")
    book.save(temp)
    os.replace(temp, EXCEL_FILE)


def write_global(state):
    timestamp = now_iso()
    with connect_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute("DELETE FROM app_state")
        conn.executemany("INSERT INTO app_state(storage_key,value_json,updated_at) VALUES(?,?,?)", [(k, v, timestamp) for k, v in state.items() if k != "infinix-accounts"])
    export_excel(state)


def owner_matches(record, user):
    return str(record.get("seller", "")).strip().lower() in {user["username"].lower(), user["email"].lower()}


def apply_state(payload, user):
    if not isinstance(payload, dict) or len(payload) > 100:
        raise ValueError("Invalid state payload")
    clean = {str(k): str(v) for k, v in payload.items() if isinstance(k, str) and k != "infinix-accounts" and len(str(v)) <= 2_000_000}
    current = read_global()
    role = user["role"]
    timestamp = now_iso()
    with connect_db() as conn:
        for key, value in clean.items():
            if key in SAVED_PRIVATE_KEYS:
                conn.execute("INSERT INTO user_state(user_id,storage_key,value_json,updated_at) VALUES(?,?,?,?) ON CONFLICT(user_id,storage_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at", (user["id"], key, value, timestamp))
    clean = {key: value for key, value in clean.items() if key not in PRIVATE_KEYS}
    if role == "admin":
        current.update(clean)
        # A localStorage write cannot create, edit, or delete server accounts.
        current.pop("infinix-accounts", None)
        write_global(current)
    else:
        with connect_db() as conn:
            if role == "shopkeeper" and "infinix-products" in clean:
                incoming = parse_json(clean["infinix-products"], [])
                existing = parse_json(current.get("infinix-products"), [])
                owner = {u.lower() for u in (user["username"], user["email"])}
                others = [p for p in existing if str(p.get("seller", "")).lower() not in owner]
                own = []
                for product in incoming:
                    if not isinstance(product, dict):
                        continue
                    if str(product.get("seller", "")).lower() in owner:
                        product["seller"] = user["username"]
                        own.append(product)
                current["infinix-products"] = json.dumps(others + own, ensure_ascii=False)
                conn.execute("INSERT INTO app_state(storage_key,value_json,updated_at) VALUES(?,?,?) ON CONFLICT(storage_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at", ("infinix-products", current["infinix-products"], timestamp))
            if role == "customer" and "infinix-orders" in clean:
                existing = parse_json(current.get("infinix-orders"), [])
                old_ids = {str(o.get("id")) for o in existing}
                incoming = parse_json(clean["infinix-orders"], [])
                for order in incoming:
                    if not isinstance(order, dict) or str(order.get("id")) in old_ids:
                        continue
                    if str(order.get("customer", "")).lower() not in {user["username"].lower(), user["email"].lower()}:
                        continue
                    if not isinstance(order.get("items"), list):
                        continue
                    existing.append(order)
                current["infinix-orders"] = json.dumps(existing, ensure_ascii=False)
                conn.execute("INSERT INTO app_state(storage_key,value_json,updated_at) VALUES(?,?,?) ON CONFLICT(storage_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at", ("infinix-orders", current["infinix-orders"], timestamp))
            if role == "customer" and "infinix-plus-requests" in clean:
                existing = parse_json(current.get("infinix-plus-requests"), [])
                existing_ids = {str(r.get("id")) for r in existing}
                for req in parse_json(clean["infinix-plus-requests"], []):
                    if not isinstance(req, dict) or str(req.get("id")) in existing_ids:
                        continue
                    if str(req.get("username", "")).lower() == user["username"].lower() and req.get("status") == "Pending":
                        existing.append(req)
                current["infinix-plus-requests"] = json.dumps(existing, ensure_ascii=False)
                conn.execute("INSERT INTO app_state(storage_key,value_json,updated_at) VALUES(?,?,?) ON CONFLICT(storage_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at", ("infinix-plus-requests", current["infinix-plus-requests"], timestamp))
            if role == "shopkeeper" and "infinix-orders" in clean:
                incoming = parse_json(clean["infinix-orders"], [])
                existing = parse_json(current.get("infinix-orders"), [])
                by_id = {str(o.get("id")): o for o in incoming if isinstance(o, dict)}
                owner = {user["username"].lower(), user["email"].lower()}
                changed = False
                for order in existing:
                    proposal = by_id.get(str(order.get("id")))
                    if not proposal or not any(str(i.get("seller", "")).lower() in owner for i in order.get("items", [])):
                        continue
                    if proposal.get("status") == "Out for delivery" and order.get("status") == "Processing":
                        order["status"] = "Out for delivery"
                        order["deliveryApproval"] = None
                        changed = True
                    approval = proposal.get("deliveryApproval")
                    if order.get("status") == "Out for delivery" and isinstance(approval, dict) and approval.get("status") == "Pending":
                        order["deliveryApproval"] = {"status": "Pending", "requestedAt": approval.get("requestedAt") or now_iso(), "seller": user["username"]}
                        changed = True
                if changed:
                    current["infinix-orders"] = json.dumps(existing, ensure_ascii=False)
                    conn.execute("INSERT INTO app_state(storage_key,value_json,updated_at) VALUES(?,?,?) ON CONFLICT(storage_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at", ("infinix-orders", current["infinix-orders"], timestamp))
        export_excel(current)


class ShopHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, status, data, cookie=None, clear_cookie=False):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        if cookie:
            secure = "; Secure" if os.environ.get("RENDER") or self.headers.get("X-Forwarded-Proto", "").lower().startswith("https") else ""
            self.send_header("Set-Cookie", f"infinix_session={cookie}; Path=/; HttpOnly; SameSite=Strict; Max-Age={SESSION_SECONDS}{secure}")
        if clear_cookie:
            self.send_header("Set-Cookie", "infinix_session=; Path=/; HttpOnly; SameSite=Strict; Max-Age=0")
        self.end_headers()
        self.wfile.write(body)

    def body_json(self):
        if self.headers.get("Content-Type", "").split(";")[0].strip().lower() != "application/json":
            raise ValueError("JSON request required")
        length = int(self.headers.get("Content-Length", "0"))
        if length < 1 or length > 2_000_000:
            raise ValueError("Invalid request size")
        return json.loads(self.rfile.read(length))

    def cookie_token(self):
        for item in self.headers.get("Cookie", "").split(";"):
            name, sep, value = item.strip().partition("=")
            if sep and name == "infinix_session":
                return value
        return ""

    def current_user(self):
        return get_user_by_session(self.cookie_token())

    def require_same_origin(self):
        origin = self.headers.get("Origin")
        if origin and urlsplit(origin).netloc.lower() != self.headers.get("Host", "").lower():
            raise PermissionError("Cross-site request blocked")

    def do_GET(self):
        path = urlsplit(self.path).path
        if path == "/api/session":
            user = self.current_user()
            self.send_json(200, {"user": user})
        elif path == "/api/state":
            self.send_json(200, get_visible_state(self.current_user()))
        else:
            name = Path(path).name.lower()
            if any(ext in name for ext in (".db", ".sqlite", ".sqlite3", ".xlsx", ".py", ".txt", ".ndjson")) or name.startswith("."):
                self.send_error(404)
                return
            super().do_GET()

    def do_HEAD(self):
        path = urlsplit(self.path).path
        name = Path(path).name.lower()
        if any(ext in name for ext in (".db", ".sqlite", ".sqlite3", ".xlsx", ".py", ".txt", ".ndjson")) or name.startswith("."):
            self.send_error(404)
            return
        super().do_HEAD()

    def do_POST(self):
        path = urlsplit(self.path).path
        try:
            self.require_same_origin()
            data = self.body_json()
            if not isinstance(data, dict):
                raise ValueError("Expected a JSON object")
            if path == "/api/register":
                username = str(data.get("username", "")).strip()
                email = str(data.get("email", "")).strip().lower()
                name = str(data.get("name", "")).strip()
                password = str(data.get("password", ""))
                role = str(data.get("role", ""))
                if len(username) < 3 or len(username) > 32 or not username.replace("_", "").isalnum():
                    raise ValueError("Username must be 3–32 letters, numbers, or underscores")
                if len(email) > 254 or "@" not in email or not name or len(name) > 100:
                    raise ValueError("Enter a valid email and name")
                if len(password) < 8 or len(password) > 200:
                    raise ValueError("Password must be at least 8 characters")
                if role not in {"customer", "shopkeeper"}:
                    raise ValueError("Choose customer or seller")
                if username.casefold() == ADMIN_USERNAME.casefold() and ADMIN_USERNAME:
                    raise ValueError("That username is reserved")
                salt, digest = password_digest(password)
                try:
                    with connect_db() as conn:
                        cur = conn.execute("INSERT INTO users(username,email,name,role,salt,password_hash,created_at) VALUES(?,?,?,?,?,?,?)", (username, email, name, role, salt, digest, now_iso()))
                        uid = cur.lastrowid
                except sqlite3.IntegrityError:
                    self.send_json(409, {"error": "That username or email is already registered"}); return
                user = public_user({"id": uid, "username": username, "email": email, "name": name, "role": role, "created_at": now_iso()})
                token = issue_session(uid)
                export_excel(read_global())
                self.send_json(201, {"user": user}, cookie=token)
            elif path == "/api/login":
                identifier = str(data.get("identifier", "")).strip()
                password = str(data.get("password", ""))
                if not identifier or not password:
                    raise ValueError("Enter your username or email and password")
                if ADMIN_USERNAME and identifier.casefold() == ADMIN_USERNAME.casefold() and ADMIN_PASSWORD and hmac.compare_digest(password, ADMIN_PASSWORD):
                    salt, digest = password_digest(ADMIN_PASSWORD)
                    with connect_db() as conn:
                        row = conn.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
                        if row:
                            uid = row["id"]
                            conn.execute("UPDATE users SET username=?,email=?,name=?,salt=?,password_hash=? WHERE id=?", (ADMIN_USERNAME, "admin@infinix.invalid", "Administrator", salt, digest, uid))
                        else:
                            uid = conn.execute("INSERT INTO users(username,email,name,role,salt,password_hash,created_at) VALUES(?,?,?,?,?,?,?)", (ADMIN_USERNAME, "admin@infinix.invalid", "Administrator", "admin", salt, digest, now_iso())).lastrowid
                    token = issue_session(uid)
                    user = get_user_by_session(token)
                    export_excel(read_global())
                    self.send_json(200, {"user": user}, cookie=token); return
                with connect_db() as conn:
                    row = conn.execute("SELECT * FROM users WHERE username=? COLLATE NOCASE OR email=? COLLATE NOCASE", (identifier, identifier)).fetchone()
                if not row or row["role"] == "admin" or not verify_password(password, row["salt"], row["password_hash"]):
                    self.send_json(401, {"error": "Invalid username or password"}); return
                token = issue_session(row["id"])
                self.send_json(200, {"user": get_user_by_session(token)}, cookie=token)
            elif path == "/api/logout":
                token = self.cookie_token()
                with connect_db() as conn:
                    conn.execute("DELETE FROM sessions WHERE token_hash=?", (token_digest(token),))
                self.send_json(200, {"ok": True}, clear_cookie=True)
            elif path == "/api/state":
                user = self.current_user()
                if not user:
                    self.send_json(401, {"error": "Sign in required"}); return
                apply_state(data, user)
                self.send_json(200, {"saved": True})
            elif path == "/api/users/delete":
                user = self.current_user()
                if not user or user["role"] != "admin":
                    self.send_json(403, {"error": "Administrator access required"}); return
                target = str(data.get("username", "")).strip()
                with connect_db() as conn:
                    victim = conn.execute("SELECT id,role FROM users WHERE username=? COLLATE NOCASE", (target,)).fetchone()
                    if not victim or victim["role"] == "admin":
                        self.send_json(404, {"error": "User not found"}); return
                    conn.execute("DELETE FROM sessions WHERE user_id=?", (victim["id"],))
                    conn.execute("DELETE FROM user_state WHERE user_id=?", (victim["id"],))
                    conn.execute("DELETE FROM users WHERE id=?", (victim["id"],))
                export_excel(read_global())
                self.send_json(200, {"deleted": True})
            else:
                self.send_json(404, {"error": "Not found"})
        except PermissionError as error:
            self.send_json(403, {"error": str(error)})
        except (ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error)})
        except Exception:
            self.send_json(500, {"error": "Server could not complete the request"})

    do_PUT = do_POST


if __name__ == "__main__":
    if len(SESSION_SECRET) < 32:
        if os.environ.get("RENDER"):
            raise RuntimeError("Set SESSION_SECRET to a random value at least 32 characters long in Render environment variables")
        SESSION_SECRET = secrets.token_urlsafe(48)
    if os.environ.get("RENDER") and (not ADMIN_USERNAME or len(ADMIN_PASSWORD) < 12):
        raise RuntimeError("Set ADMIN_USERNAME and an ADMIN_PASSWORD of at least 12 characters in Render environment variables")
    connect_db().close()
    export_excel(read_global())
    port = int(os.environ.get("PORT", "8000"))
    ThreadingHTTPServer(("0.0.0.0", port), ShopHandler).serve_forever()
